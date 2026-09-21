from pathlib import Path

import openpyxl

from config_mapeo import cargar_catalogo_formularios, cargar_config, construir_indice_alias
from encabezados import detectar_encabezados
from mapeo import construir_fila_salida
from nombre_hoja import generar_nombre_hoja
from plantilla import cargar_variables_plantilla


NOMBRE_HOJA_GENERAL = "Consolidado General"


def _nombre_unico(nombre, usados):
    if nombre not in usados:
        usados.add(nombre)
        return nombre
    n = 2
    while True:
        candidato = f"{nombre[:28]} ({n})"
        if candidato not in usados:
            usados.add(candidato)
            return candidato
        n += 1


def procesar_archivos(rutas_entrada, on_evento=None):
    """Consolida los excels de rutas_entrada en un libro en memoria (no lo guarda).

    on_evento(mensaje, tipo), si se pasa, se llama en tiempo real con el avance
    (tipo: "info", "advertencia" o "exito").

    Devuelve (libro_salida, resultado) donde libro_salida es el Workbook de
    openpyxl (para guardarlo despues con libro_salida.save(ruta) una vez el
    usuario elija donde) y resultado es un dict:
      {
        "total_filas": int,
        "total_duplicados": int,
        "advertencias": [str, ...],
        "archivos": [
          {"archivo": str, "hoja": str, "filas": int, "duplicados_dni": [str, ...]},
          ...
        ],
      }
    """

    def emitir(mensaje, tipo="info"):
        if on_evento:
            on_evento(mensaje, tipo)

    variables_plantilla = cargar_variables_plantilla()
    indice_alias = construir_indice_alias(cargar_config())
    catalogo_formularios = cargar_catalogo_formularios()
    indice_dni = variables_plantilla.index("DNI")

    wb_salida = openpyxl.Workbook()
    wb_salida.remove(wb_salida.active)
    nombres_usados = {NOMBRE_HOJA_GENERAL}
    advertencias = []
    archivos_resultado = []
    filas_generales = []
    total_filas = 0
    total_duplicados = 0

    for ruta in rutas_entrada:
        ruta = Path(ruta)
        emitir(f"Procesando {ruta.name}...")
        wb_entrada = openpyxl.load_workbook(ruta, data_only=True)
        ws = wb_entrada.active

        fila_encabezado, mapeo = detectar_encabezados(ws, indice_alias)
        if not mapeo:
            mensaje = f"{ruta.name}: no se reconocio ningun encabezado, se omitio."
            advertencias.append(mensaje)
            emitir(mensaje, "advertencia")
            wb_entrada.close()
            continue

        nombre_base, reconocido = generar_nombre_hoja(ruta.name, catalogo_formularios)
        if not reconocido:
            mensaje = f"{ruta.name}: nombre de formulario no catalogado, se uso '{nombre_base}'."
            advertencias.append(mensaje)
            emitir(mensaje, "advertencia")
        nombre_hoja = _nombre_unico(nombre_base, nombres_usados)

        ws_salida = wb_salida.create_sheet(title=nombre_hoja)
        ws_salida.append(variables_plantilla)

        dnis_vistos = set()
        dnis_duplicados = []
        filas_agregadas = 0
        for fila in ws.iter_rows(min_row=fila_encabezado + 1):
            if all(celda.value is None for celda in fila):
                continue
            valores_fila = {celda.column: celda.value for celda in fila}
            fila_salida = construir_fila_salida(valores_fila, mapeo, variables_plantilla, nombre_hoja)

            dni = fila_salida[indice_dni]
            if dni is not None and str(dni).strip() != "":
                clave_dni = str(dni).strip()
                if clave_dni in dnis_vistos:
                    dnis_duplicados.append(clave_dni)
                    continue
                dnis_vistos.add(clave_dni)

            ws_salida.append(fila_salida)
            filas_generales.append(fila_salida)
            filas_agregadas += 1

        if dnis_duplicados:
            mensaje = (
                f"{nombre_hoja}: se descartaron {len(dnis_duplicados)} fila(s) con DNI "
                "duplicado dentro de la misma hoja (se conservo el primer registro)."
            )
            advertencias.append(mensaje)
            emitir(mensaje, "advertencia")

        emitir(f"-> hoja '{nombre_hoja}': {filas_agregadas} fila(s) agregadas.")
        archivos_resultado.append(
            {
                "archivo": ruta.name,
                "hoja": nombre_hoja,
                "filas": filas_agregadas,
                "duplicados_dni": dnis_duplicados,
            }
        )
        total_filas += filas_agregadas
        total_duplicados += len(dnis_duplicados)
        wb_entrada.close()

    ws_general = wb_salida.create_sheet(title=NOMBRE_HOJA_GENERAL, index=0)
    ws_general.append(variables_plantilla)
    for fila_salida in filas_generales:
        ws_general.append(fila_salida)
    emitir(f"-> hoja '{NOMBRE_HOJA_GENERAL}': {len(filas_generales)} fila(s) agregadas.")

    emitir("Consolidado generado en memoria. Elige donde guardarlo.", "exito")

    resultado = {
        "total_filas": total_filas,
        "total_duplicados": total_duplicados,
        "advertencias": advertencias,
        "archivos": archivos_resultado,
    }
    return wb_salida, resultado
